# -*- coding: utf-8 -*-
"""
Usability survey analysis for the Real-Time Child Monitoring System.

Usage:  python analyze_surveys.py        (run from the evaluation/ folder)

Reads the three response files in data/, converts every answer to the 1-5
Likert scale, and writes results.md with per-item and per-group statistics.
Requires openpyxl.

Three separate questionnaires were used, one per role, each with its own
wording. Item numbers are therefore NOT comparable across files.

The surveys were administered at two child development centers in Guayaquil,
Ecuador. The questions in the header row have been translated into English, but
the answers are still the Spanish strings the forms recorded, because they are
the raw data; data/questions.md has every item in both languages.

Answers are compared in lower case and without accents, because capitalisation
and accents vary between files ("Muy Facil" and "Muy facil" are the same
answer).

The five response scales, all five-point Likert:

  usefulness    Nada util 1 · Poco util 2 · Neutral 3 · Util 4 · Muy util 5
  agreement     Definitivamente no 1 · Probablemente no 2 · Tal vez 3 ·
                Probablemente si 4 · Definitivamente si 5
  satisfaction  Muy insatisfecho 1 ... Muy satisfecho 5
  ease of use   Muy dificil 1 · Dificil 2 · Neutral 3 · Facil 4 · Muy facil 5
  quality       Mala 1 · Regular 2 · Buena 3 · Muy buena 4 · Excelente 5

Standard deviations are sample standard deviations (n-1).
"""

import io
import statistics
from collections import Counter

import openpyxl

OUTPUT = "results.md"

FILES = [
    ("parents", "data/parents_responses.xlsx"),
    ("tutors", "data/tutors_responses.xlsx"),
    ("coordinators", "data/coordinators_responses.xlsx"),
]

SCALE = {
    # usefulness
    "nada util": 1, "poco util": 2, "neutral": 3, "util": 4, "muy util": 5,
    # agreement
    "definitivamente no": 1, "probablemente no": 2, "tal vez": 3,
    "probablemente si": 4, "definitivamente si": 5,
    # satisfaction
    "muy insatisfecho": 1, "insatisfecho": 2,
    "ni satisfecho ni insatisfecho": 3, "satisfecho": 4, "muy satisfecho": 5,
    # ease of use
    "muy dificil": 1, "dificil": 2, "facil": 4, "muy facil": 5,
    # perceived quality
    "mala": 1, "regular": 2, "buena": 3, "muy buena": 4, "excelente": 5,
}

ACCENTS = {"á": "a", "é": "e", "í": "i", "ó": "o", "ú": "u", "ñ": "n"}

# The five scales in full, for the report: value, English wording, and the
# Spanish option as the form recorded it.
#
# The two lowest points of the ease-of-use scale are marked as inferred: no
# respondent chose them, and the option list of that question was not kept, so
# their exact wording is a reconstruction.
SCALES = [
    ("usefulness", [
        (1, "Not at all useful", "Nada útil"),
        (2, "Slightly useful", "Poco útil"),
        (3, "Neutral", "Neutral"),
        (4, "Useful", "Útil"),
        (5, "Very useful", "Muy útil")]),
    # Five of the seven items on this scale ask whether something is
    # sufficient or whether it was provided, not what the respondent
    # intends to do, so it is named for agreement rather than intention.
    ("agreement", [
        (1, "Definitely not", "Definitivamente no"),
        (2, "Probably not", "Probablemente no"),
        (3, "Maybe", "Tal vez"),
        (4, "Probably yes", "Probablemente sí"),
        (5, "Definitely yes", "Definitivamente sí")]),
    # The midpoint was documented as "Ni satisfecho ni insatisfecho" but the
    # form stored it as "Neutral", which is what the file contains.
    ("satisfaction", [
        (1, "Very dissatisfied", "Muy insatisfecho"),
        (2, "Dissatisfied", "Insatisfecho"),
        (3, "Neutral", "Neutral"),
        (4, "Satisfied", "Satisfecho"),
        (5, "Very satisfied", "Muy satisfecho")]),
    ("ease of use", [
        (1, "Very difficult", "Muy difícil *(inferred)*"),
        (2, "Difficult", "Difícil *(inferred)*"),
        (3, "Neutral", "Neutral"),
        (4, "Easy", "Fácil"),
        (5, "Very easy", "Muy fácil")]),
    ("quality", [
        (1, "Bad", "Mala"),
        (2, "Fair", "Regular"),
        (3, "Good", "Buena"),
        (4, "Very good", "Muy buena"),
        (5, "Excellent", "Excelente")]),
]


def normalise(value):
    text = str(value).strip().lower()
    for accented, plain in ACCENTS.items():
        text = text.replace(accented, plain)
    return text


def spell(number):
    """Small numbers read better as words in prose, and must stay derived."""
    words = ["zero", "one", "two", "three", "four", "five",
             "six", "seven", "eight", "nine", "ten"]
    return words[number] if number < len(words) else str(number)


def scale_name(answers):
    """Identify the scale from the answers given, not from the question."""
    given = set(answers)
    if given & {"muy util", "util", "poco util", "nada util"}:
        return "usefulness"
    if given & {"muy facil", "facil", "dificil", "muy dificil"}:
        return "ease of use"
    if given & {"definitivamente si", "probablemente si", "tal vez"}:
        return "agreement"
    if given & {"muy satisfecho", "satisfecho", "insatisfecho"}:
        return "satisfaction"
    if given & {"excelente", "muy buena", "buena", "regular", "mala"}:
        return "quality"
    return "unknown"


lines = []
per_group = {}
sizes = {}
# How many times each option was picked, keyed by (scale, option). It has to
# be keyed by the scale too: "Neutral" is the midpoint of three different
# scales, and counting it once mixes them into a single wrong number.
chosen = Counter()

# (scale, value) -> English wording, for the per-item breakdown
ENGLISH = {(name, value): english
           for name, points in SCALES for value, english, _ in points}

lines.append("# Survey results")
lines.append("")
lines.append("Generated by `analyze_surveys.py` from the response files in")
lines.append("`data/`. **Do not edit by hand**: re-run the script instead.")
lines.append("")
lines.append("Every item was answered on a five-point Likert scale. Standard "
             "deviations are sample standard deviations (n-1).")
lines.append("")
lines.append("*Distinct response patterns* counts how many of the rows "
             "differ from each other. With a handful of options and seven or "
             "eight questions, identical rows are expected.")

sections = []

for group, path in FILES:
    book = openpyxl.load_workbook(path, data_only=True)
    rows = list(book[book.sheetnames[0]].iter_rows(values_only=True))
    questions, answers = rows[0], [r for r in rows[1:]
                                   if any(c is not None for c in r)]
    sizes[group] = len(answers)

    file_warnings = []
    if len(book.sheetnames) > 1:
        file_warnings.append(
            "**This file has more than one sheet** (%s). Only the first is "
            "read." % ", ".join(book.sheetnames))
    if len(answers) < 2:
        raise SystemExit(
            "At least two response rows are needed to compute a standard "
            "deviation, and %s has %d." % (path, len(answers)))

    section = []
    section.append("")
    section.append("## %s (n = %d)" % (group.capitalize(), len(answers)))
    section.append("")
    section.append("Source: `%s`" % path)
    section.append("")
    section.append("| # | Question | Scale | Answers chosen | Mean | Median | SD | 4 or 5 |")
    section.append("|---|---|---|---|---|---|---|---|")

    by_item = {}
    warnings = list(file_warnings)
    for index, question in enumerate(questions):
        column = [normalise(row[index]) for row in answers
                  if row[index] is not None]

        # A blank cell would otherwise shrink an item's n with nothing to show
        # for it: the table has no n column, so it would be invisible.
        if len(column) < len(answers):
            warnings.append(
                "**Item %d was left blank by %d of the %d respondents**, so "
                "its statistics are over %d answers."
                % (index + 1, len(answers) - len(column), len(answers),
                   len(column)))

        unmapped = sorted({c for c in column if c not in SCALE})
        if unmapped:
            warnings.append(
                "**Item %d has answers outside the known scales:** %s. The "
                "whole item is left out of the table and of the group "
                "statistics below." % (index + 1, ", ".join(unmapped)))
            continue

        this = scale_name(column)
        if this == "unknown":
            # Can happen when every answer is a label shared by several scales,
            # such as "Neutral": there is then no way to tell which scale it is.
            warnings.append(
                "**Item %d cannot be assigned to a scale** from the answers it "
                "received (%s), so it is left out of the table and of the "
                "group statistics below."
                % (index + 1, ", ".join(sorted(set(column)))))
            continue

        values = [SCALE[c] for c in column]
        by_item[index + 1] = (values, question)
        chosen.update((this, c) for c in column)
        per_option = Counter(values)
        # each option with its English wording, not just the number
        detalle = "<br>".join(
            "%d × **%s** (%d)" % (per_option[v], ENGLISH[(this, v)], v)
            for v in sorted(per_option, reverse=True))
        top_two = 100.0 * sum(1 for v in values if v >= 4) / len(values)
        section.append("| %d | %s | [%s](#%s) | %s | **%.2f** | %.1f | %.2f | %.0f %% |"
                       % (index + 1, str(question).strip(), this,
                          this.replace(" ", "-"), detalle,
                          statistics.mean(values), statistics.median(values),
                          statistics.stdev(values), top_two))

    for warning in warnings:
        section.append("")
        section.append(warning)

    flat = [v for values, _ in by_item.values() for v in values]
    per_group[group] = flat
    per_value = Counter(flat)
    section.append("")
    section.append("**Group mean %.2f** (SD %.2f) over %d ratings. "
                   "Distribution: %s. No rating below %d."
                   % (statistics.mean(flat), statistics.stdev(flat), len(flat),
                      ", ".join("%d × %d" % (per_value[v], v)
                                for v in sorted(per_value, reverse=True)),
                      min(flat)))
    # Blanks are kept as None rather than dropped: dropping them could make two
    # different rows look identical.
    patterns = [tuple(normalise(c) if c is not None else None for c in row)
                for row in answers]
    section.append("")
    section.append("Distinct response patterns: %d of %d."
                   % (len(set(patterns)), len(answers)))
    sections.append(section)

flat = [v for values in per_group.values() for v in values]

# The summary goes first, before the per-group sections.
lines.append("")
lines.append("## All three groups")
lines.append("")
lines.append("| Group | n | Ratings | Mean | SD | 4 or 5 |")
lines.append("|---|---|---|---|---|---|")
for group, values in per_group.items():
    lines.append("| %s | %d | %d | **%.2f** | %.2f | %.0f %% |"
                 % (group.capitalize(), sizes[group], len(values),
                    statistics.mean(values), statistics.stdev(values),
                    100.0 * sum(1 for v in values if v >= 4) / len(values)))
lines.append("| **All** | **%d** | **%d** | **%.2f** | **%.2f** | **%.0f %%** |"
             % (sum(sizes.values()), len(flat), statistics.mean(flat),
                statistics.stdev(flat),
                100.0 * sum(1 for v in flat if v >= 4) / len(flat)))
lines.append("")
lines.append("**No rating in any questionnaire fell below the neutral "
             "midpoint**: the minimum of all %d ratings is %d. This is "
             "consistent with courtesy bias in a demonstration setting and is "
             "reported as a limitation in the paper." % (len(flat), min(flat)))

lines.append("")
lines.append("## Response scales")
lines.append("")
lines.append("%s scales were used, all five-point Likert. Each item below is "
             "labelled with the one it used. **Which scale an item used is "
             "determined by the answers it received**, not by the wording of "
             "the question: the location items ask how the function is "
             "*perceived* but were answered on the ease-of-use scale, so they "
             "are reported as perceived ease of use." % spell(len(SCALES)).capitalize())
lines.append("")
lines.append("The *As recorded* column is the Spanish option stored in the "
             "files; *Times chosen* counts it across all %s questionnaires."
             % spell(len(FILES)))

for name, points in SCALES:
    lines.append("")
    lines.append("### %s" % name.capitalize())
    lines.append("")
    lines.append("| Value | English | As recorded | Times chosen |")
    lines.append("|---|---|---|---|")
    for value, english, spanish in points:
        picked = chosen.get(
            (name, normalise(spanish.replace(" *(inferred)*", ""))), 0)
        lines.append("| %d | %s | %s | %s |"
                     % (value, english, spanish, picked if picked else "—"))

lines.append("")
por_escala = {name: sum(n for (s, _), n in chosen.items() if s == name)
              for name, _ in SCALES}
lines.append("The %s scales together account for all %d ratings: %s."
             % (spell(len(SCALES)), sum(por_escala.values()),
                ", ".join("%s %d" % (name, por_escala[name])
                          for name, _ in SCALES)))
if sum(por_escala.values()) != len(flat):
    lines.append("")
    lines.append("**Warning: the scales account for %d ratings but %d were "
                 "counted.**" % (sum(por_escala.values()), len(flat)))
lines.append("")
lines.append("**No option below %d was ever chosen**, on any scale, by any of "
             "the %d participants. The two lowest points of the ease-of-use "
             "scale are marked *inferred*: nobody chose them and the option "
             "list of that question was not kept, so their exact wording is a "
             "reconstruction. An option nobody chose contributes no rating, so "
             "nothing in this report depends on it."
             % (min(flat), sum(sizes.values())))

for section in sections:
    lines.extend(section)

io.open(OUTPUT, "w", encoding="utf-8").write("\n".join(lines) + "\n")
print("written", OUTPUT)
